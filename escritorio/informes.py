# -*- coding: utf-8 -*-
"""
Informe para imprimir, en Excel y en PDF.

Lo justo: por persona, los dias trabajados, cuanto se le pago y cuanto se le
debe; y despues el detalle de cada dia. Las horas van como reloj (10:42), no
con decimales.
"""

from nucleo import pesos, hhmm_txt, nombre_dia

ROJO = "B4141F"
VERDE = "16793B"
AMBAR = "8A5A00"
TINTA = "15100F"
GRIS = "EDE9E4"
AUTOR = "Generado con Control de Horas, desarrollado por Macoem."


def _fecha(iso):
    return "%s %s-%s-%s" % (nombre_dia(iso)[:3], iso[8:10], iso[5:7], iso[:4])


def _estado(d):
    """Como esta pagado ese dia: todo, solo una parte, o nada."""
    if not d.get("completa", True):
        return "falta marcar"
    if d.get("pagado"):
        return "pagado el %s-%s" % (d["pagado_en"][8:10], d["pagado_en"][5:7])
    if d.get("pagado_normal"):
        return "pagadas solo las normales"
    if d.get("pagado_extra"):
        return "pagadas solo las extra"
    return "por pagar"


def _notas(r):
    return [
        "Horas trabajadas = de la entrada a la salida, menos la colacion.",
        "Horas extra: lo que pasa de %s horas en el dia. La hora extra se paga con %s."
        % (r["contrato"], r["regla_extra"]),
        "Las horas normales y las extra se pagan por separado: un dia puede tener "
        "una parte pagada y la otra pendiente.",
        "Lo pagado es el monto que quedo registrado al marcar cada parte como pagada.",
        AUTOR,
    ]


# --------------------------------------------------------------------- Excel
def a_excel(r, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    borde = Border(*[Side(style="thin", color="D5CFC9")] * 4)
    cab = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=TINTA)

    def encabezado(hoja, fila, columnas):
        for i, (nombre, ancho) in enumerate(columnas, start=1):
            c = hoja.cell(row=fila, column=i, value=nombre)
            c.font = cab; c.fill = relleno; c.border = borde
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            hoja.column_dimensions[get_column_letter(i)].width = ancho
        hoja.row_dimensions[fila].height = 28

    # ---- Resumen por persona
    h = wb.active
    h.title = "Pagos por persona"
    h["A1"] = r["negocio"]; h["A1"].font = Font(size=16, bold=True, color=ROJO)
    h["A2"] = "%s  |  Pagos  |  %s" % (r["ciudad"], r["titulo"])
    h["A2"].font = Font(size=11, bold=True, color=TINTA)
    fila = 4
    encabezado(h, fila, [("Trabajador", 26), ("Dias", 8), ("Pagado en normales", 18),
                         ("Pagado en extra", 16), ("Debe en normales", 17),
                         ("Debe en extra", 15), ("SE LE DEBE", 16)])
    for f in r["filas"]:
        fila += 1
        for i, v in enumerate([f["nombre"], f["turnos"], f["pagado_normal"],
                               f["pagado_extra"], f["por_pagar_normal"],
                               f["por_pagar_extra"], f["por_pagar"]], start=1):
            c = h.cell(row=fila, column=i, value=v)
            c.border = borde
            if i >= 3:
                c.number_format = '"$"#,##0'
            if i == 7 and f["por_pagar"]:
                c.font = Font(bold=True, color=ROJO)
    fila += 2
    for linea in _notas(r):
        h.cell(row=fila, column=1, value=linea).font = Font(size=9, color="6C625C")
        fila += 1
    try:
        import base64, io
        import imagen_marca
        from openpyxl.drawing.image import Image as ImagenXL
        logo = ImagenXL(io.BytesIO(base64.b64decode(imagen_marca.PDF_JPEG)))
        logo.width = logo.height = 66
        h.add_image(logo, "G1")                 # arriba a la derecha
        h.row_dimensions[1].height = 24
        h.row_dimensions[2].height = 24
    except Exception:
        pass                                     # sin logo, el informe igual sirve
    h.page_setup.orientation = "landscape"
    h.page_setup.fitToWidth = 1
    h.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- Detalle de dias
    d = wb.create_sheet("Dias")
    d["A1"] = "Dias trabajados  -  %s" % r["titulo"]
    d["A1"].font = Font(size=13, bold=True, color=TINTA)
    fila = 3
    encabezado(d, fila, [("Trabajador", 22), ("Fecha", 16), ("Horas", 10),
                         ("H. normales", 13), ("$ normales", 14), ("H. extra", 11),
                         ("$ extra", 13), ("Total del dia", 14), ("Estado", 24)])
    for f in r["filas"]:
        for x in sorted(f["dias"], key=lambda y: y["fecha"]):
            fila += 1
            completo = x.get("completa", True)
            valores = [f["nombre"], _fecha(x["fecha"]),
                       hhmm_txt(x["horas"]) if completo else "-",
                       hhmm_txt(x["normales"]) if completo else "-",
                       x["pago_normal"] if completo else 0,
                       hhmm_txt(x["extra"]) if x["extra"] else "-",
                       x["pago_extra"] if completo else 0,
                       x["total"] if completo else 0, _estado(x)]
            for i, v in enumerate(valores, start=1):
                c = d.cell(row=fila, column=i, value=v)
                c.border = borde
                if i in (3, 4, 6):
                    c.alignment = Alignment(horizontal="center")
                if i in (5, 7, 8):
                    c.number_format = '"$"#,##0'
                if i == 9:
                    c.font = Font(bold=True, color=VERDE if x.get("pagado") else
                                  (AMBAR if not completo else ROJO))
    d.freeze_panes = "A4"
    d.print_title_rows = "3:3"
    d.page_setup.fitToWidth = 1
    d.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------- PDF
def _logo_pdf(lado):
    """
    El logo del local para la cabecera del PDF. Va en JPEG porque reportlab lo
    incrusta sin necesitar Pillow, que no va dentro del programa.
    """
    try:
        import base64, io
        import imagen_marca
        from reportlab.platypus import Image
        return Image(io.BytesIO(base64.b64decode(imagen_marca.PDF_JPEG)),
                     width=lado, height=lado)
    except Exception:
        return None


def a_pdf(r, ruta):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    rojo, verde, ambar = (colors.HexColor("#" + x) for x in (ROJO, VERDE, AMBAR))
    tinta, gris, suave = colors.HexColor("#" + TINTA), colors.HexColor("#" + GRIS), \
        colors.HexColor("#6C625C")
    est = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=est["Title"], fontSize=17, textColor=rojo,
                        alignment=0, spaceAfter=1)
    h2 = ParagraphStyle("h2", parent=est["Normal"], fontSize=11, textColor=tinta,
                        fontName="Helvetica-Bold", spaceAfter=10)
    sec = ParagraphStyle("sec", parent=est["Normal"], fontSize=11, textColor=tinta,
                         fontName="Helvetica-Bold", spaceBefore=14, spaceAfter=5)
    nota = ParagraphStyle("nota", parent=est["Normal"], fontSize=7.5, textColor=suave,
                          leading=10)

    doc = SimpleDocTemplate(ruta, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Pagos - " + r["titulo"])
    titulo = [Paragraph(r["negocio"], h1),
              Paragraph("%s &nbsp;|&nbsp; Pagos &nbsp;|&nbsp; %s" % (r["ciudad"], r["titulo"]), h2)]
    logo = _logo_pdf(24 * mm)
    if logo is not None:
        cab = Table([[logo, titulo]], colWidths=[29 * mm, None])
        cab.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        cab.hAlign = "LEFT"
        hist = [cab]
    else:
        hist = titulo

    datos = [["Trabajador", "Dias", "Pagado en\nnormales", "Pagado en\nextra",
              "Debe en\nnormales", "Debe en\nextra", "SE LE\nDEBE"]]
    for f in r["filas"]:
        datos.append([f["nombre"], str(f["turnos"]), pesos(f["pagado_normal"]),
                      pesos(f["pagado_extra"]), pesos(f["por_pagar_normal"]),
                      pesos(f["por_pagar_extra"]), pesos(f["por_pagar"])])
    tabla = Table(datos, colWidths=[43 * mm, 13 * mm, 25 * mm, 23 * mm, 25 * mm,
                                    23 * mm, 27 * mm], repeatRows=1)
    tabla.hAlign = "LEFT"
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), tinta), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5CFC9")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for i, f in enumerate(r["filas"], start=1):
        if f["por_pagar"]:
            estilo += [("TEXTCOLOR", (6, i), (6, i), rojo),
                       ("FONTNAME", (6, i), (6, i), "Helvetica-Bold")]
    tabla.setStyle(TableStyle(estilo))
    hist.append(tabla)

    for f in r["filas"]:
        hist.append(Paragraph("%s &nbsp;&mdash;&nbsp; pagado %s, se le debe %s"
                              % (f["nombre"], pesos(f["pagado"]), pesos(f["por_pagar"])), sec))
        det = [["Fecha", "Horas", "Normales", "$ normales", "Extra", "$ extra",
                "Total dia", "Estado"]]
        colores = []
        for j, x in enumerate(sorted(f["dias"], key=lambda y: y["fecha"]), start=1):
            completo = x.get("completa", True)
            det.append([_fecha(x["fecha"]),
                        hhmm_txt(x["horas"]) if completo else "-",
                        hhmm_txt(x["normales"]) if completo else "-",
                        pesos(x["pago_normal"]) if completo else "-",
                        hhmm_txt(x["extra"]) if x["extra"] else "-",
                        pesos(x["pago_extra"]) if x["extra"] else "-",
                        pesos(x["total"]) if completo else "-", _estado(x)])
            colores.append(("TEXTCOLOR", (7, j), (7, j),
                            verde if x.get("pagado") else (ambar if not completo else rojo)))
        td = Table(det, colWidths=[26 * mm, 15 * mm, 18 * mm, 22 * mm, 15 * mm,
                                   20 * mm, 22 * mm, 41 * mm], repeatRows=1)
        td.hAlign = "LEFT"
        td.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), gris), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("ALIGN", (1, 0), (6, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D5CFC9")),
            ("FONTNAME", (7, 1), (7, -1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ] + colores))
        hist.append(td)

    hist.append(Spacer(1, 9 * mm))
    for linea in _notas(r):
        hist.append(Paragraph("&bull; " + linea, nota))
    doc.build(hist)
    return ruta
